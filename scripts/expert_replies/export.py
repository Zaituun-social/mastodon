"""Writes queue.json out as a human-friendly spreadsheet for reviewing
candidate posts (and near-duplicate flags) outside the terminal -- e.g. to
hand to someone who isn't going to read raw JSON.

Regenerated automatically after fetch and after each reply, and can also
be run standalone via `python -m scripts.expert_replies export`.
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import config

HEADER_FILL = PatternFill(start_color="FFD9D9D9", end_color="FFD9D9D9", fill_type="solid")
DUPLICATE_FILL = PatternFill(start_color="FFFFF2CC", end_color="FFFFF2CC", fill_type="solid")
REPLIED_FILL = PatternFill(start_color="FFD9EAD3", end_color="FFD9EAD3", fill_type="solid")

REPLY_COLUMN_HEADER = "Reply"
STATUS_ID_COLUMN_HEADER = "Status ID"

# key is None for the Reply column: it's an input field with no matching
# queue.json field, so it always renders blank on (re)export.
COLUMNS = [
    ("interest", "Interest", 22),
    ("matched_interests", "Matched Interests", 30),
    ("tag", "Tag", 18),
    ("account_acct", "From", 18),
    ("content_text", "Post", 60),
    (None, REPLY_COLUMN_HEADER, 60),
    ("is_duplicate", "Duplicate?", 12),
    ("duplicate_of", "Duplicate Of", 14),
    ("replied", "Replied?", 10),
    ("url", "URL", 40),
    ("created_at", "Created At", 20),
    ("id", STATUS_ID_COLUMN_HEADER, 14),
]


def write_xlsx(queue, path=None):
    path = path or config.CANDIDATES_XLSX_FILE

    wb = Workbook()
    ws = wb.active
    ws.title = "Candidate Posts"
    ws.sheet_view.rightToLeft = True

    for col_idx, (_, label, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    for row_idx, item in enumerate(queue, start=2):
        for col_idx, (key, _, _) in enumerate(COLUMNS, start=1):
            value = item.get(key)
            if key in ("is_duplicate", "replied"):
                value = "Yes" if value else "No"
            elif key == "matched_interests":
                value = ", ".join(value or [item.get("interest")])
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(wrap_text=True, vertical="top", horizontal="right")

        row_fill = None
        if item.get("replied"):
            row_fill = REPLIED_FILL
        elif item.get("is_duplicate"):
            row_fill = DUPLICATE_FILL

        if row_fill:
            for col_idx in range(1, len(COLUMNS) + 1):
                ws.cell(row=row_idx, column=col_idx).fill = row_fill

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    try:
        wb.save(path)
    except PermissionError:
        print(f"  could not write {path} (probably open in Excel/Numbers -- close it and re-export)")
        return None

    return path
